from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLEU = RGBColor(0x1F, 0x4E, 0x79)


def nouveau_document():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    for niveau, taille in ((1, 16), (2, 13)):
        h = doc.styles[f"Heading {niveau}"]
        h.font.name = "Calibri"
        h.font.size = Pt(taille)
        h.font.color.rgb = BLEU
    return doc


def _texte(p, texte, taille=None):
    """Ajoute du texte ; les passages entre ** sont en gras."""
    for i, morceau in enumerate(texte.split("**")):
        r = p.add_run(morceau)
        r.bold = (i % 2 == 1)
        if taille:
            r.font.size = Pt(taille)


def titre(doc, texte, niveau=1):
    doc.add_heading(texte, level=niveau)


def paragraphe(doc, texte):
    _texte(doc.add_paragraph(), texte)


def puces(doc, lignes):
    for ligne in lignes:
        _texte(doc.add_paragraph(style="List Bullet"), ligne)


def _ombrer(cellule, couleur):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    cellule._tc.get_or_add_tcPr().append(shd)


def tableau(doc, entete, lignes, largeurs=None, taille=10):
    t = doc.add_table(rows=1, cols=len(entete))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Répéter l'en-tête et interdire la coupure
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    trPr.append(OxmlElement('w:cantSplit'))

    for j, texte in enumerate(entete):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(texte)
        r.bold = True
        r.font.size = Pt(taille)
        r.font.color.rgb = RGBColor(255, 255, 255)
        _ombrer(c, "1F4E79")

    for k, ligne in enumerate(lignes):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cellules = row.cells
        for j, texte in enumerate(ligne):
            _texte(cellules[j].paragraphs[0], str(texte), taille=taille)
            if k % 2 == 1:
                _ombrer(cellules[j], "EAF1F8")

    if largeurs:
        for row in t.rows:
            for j, w in enumerate(largeurs):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()



def encadre(doc, texte, couleur="DEEAF6"):
    """Encadré coloré pour les messages clés."""
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _ombrer(c, couleur)
    _texte(c.paragraphs[0], texte)
    doc.add_paragraph()


if __name__ == "__main__":
    # Document de test pour vérifier la mise en page
    doc = nouveau_document()
    titre(doc, "Titre de section")
    titre(doc, "Sous-titre", 2)
    paragraphe(doc, "Un paragraphe avec un passage **en gras**.")
    puces(doc, ["Premier point", "Deuxième point **important**"])
    tableau(doc, ["Main", "Annonce", "Réussite"],
            [["V + 2 petits + as", "80", "83 %"], ["V 9 + 2 petits", "100", "83 %"]])
    encadre(doc, "**À retenir** : un message clé dans un encadré.")
    doc.save("test_outils.docx")
    print("test_outils.docx créé")
