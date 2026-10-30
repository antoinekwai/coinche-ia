from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from rapport_outils import titre, paragraphe, puces, tableau, encadre, BLEU


def page_titre(doc):
    for _ in range(6):
        doc.add_paragraph()

    def ligne(texte, taille, gras=False, italique=False, couleur=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texte)
        r.font.size = Pt(taille)
        r.bold, r.italic = gras, italique
        if couleur:
            r.font.color.rgb = couleur

    ligne("La coinche au banc d'essai", 30, gras=True, couleur=BLEU)
    ligne("Ce que des centaines de milliers de donnes simulées nous apprennent "
          "sur les annonces, la coinche et le jeu de la carte", 14, italique=True)
    for _ in range(8):
        doc.add_paragraph()
    ligne("Kwai", 14, gras=True)
    ligne("Synthèse à l'attention des joueurs de coinche", 11)
    doc.add_page_break()

def section_en_bref(doc):
    titre(doc, "1. En bref")
    paragraphe(doc, "Nous avons programmé un simulateur de belote coinchée, fait jouer des joueurs "
               "virtuels de différents niveaux, et comparé des façons d'annoncer, de coincher et de "
               "jouer. Chaque comparaison rejoue **les mêmes cartes en inversant les équipes**, ce qui "
               "élimine la chance de la donne. Ces enseignements valent pour le barème « on marque "
               "ou on chute la valeur du contrat ».")

    titre(doc, "Les découvertes principales", 2)
    puces(doc, [
        "**Ouvrir au palier que l'on réussit environ 4 fois sur 5.** Plus audacieux, les chutes et "
        "les coinches coûtent trop cher ; plus prudent, on laisse filer des donnes.",
        "**Chaque « bonne carte » vaut 10 à 12 points** : un as, un atout de plus, le 9 avec le valet.",
        "**Le valet long est une vraie force** (valet quatrième + 1 as : 100), **V 9 secs s'ouvrent "
        "dès qu'on a des as**, et **le 9 sans le valet reste faible**.",
        "**Le partenaire court à l'atout reste prudent** : sur un 80, avec 0 ou 1 atout, il faut "
        "2 as pour remonter.",
        "**Quand les adversaires parlent : niveau de l'équipe + 20 à 30, la marge ne faisant jamais "
        "dépasser 130.** Cette règle fait aussi bien qu'une simulation informatique.",
        "**Coincher sur les atouts** (valet, longueur, roi ou dame qui empêche la belote adverse). "
        "Plusieurs variantes plus élaborées ont été testées : aucune ne fait mieux.",
        "**Annonces et jeu de la carte vont de pair** : plus on annonce au plus juste, plus le jeu "
        "de la carte devient décisif.",
        "**Au jeu de la carte** : avec 3 atouts, on rejoue atout quand le partenaire a pris ; sur le "
        "pli du partenaire, on charge son 10 en défense et on le garde en attaque ; quand l'As "
        "gagne le pli, on le met (l'impasse ne rapporte pas en moyenne).",
        "**Écouter les enchères, y compris les passes,** aide aussi à placer les cartes cachées.",
        "**Pour le capot, compter ses plis ne suffit pas** : la forme de la main et les "
        "communications avec le partenaire comptent autant que les cartes maîtresses.",
    ])

    titre(doc, "Ce que rapportent les améliorations", 2)
    tableau(doc, ["Amélioration", "Gain par donne", "Ordre de grandeur par partie"], [
        ["Système corrigé (contre le système d'origine)", "+7", "≈ +80 points"],
        ["Ouvertures alignées sur les tableaux de mains", "+1,5", "≈ +15 points"],
        ["Réponse prudente du partenaire sur 80", "+2", "≈ +20 points"],
        ["Coinche basée sur les atouts", "+9", "≈ +100 points"],
        ["Bonne marge dans les enchères disputées (contre aucune marge)", "+8 à +20",
         "+100 à +200 points"],
        ["Meilleur jeu de la carte (selon la justesse des annonces)", "+4 à +24",
         "+45 à +260 points"],
        ["Tenir compte des enchères pour jouer la carte", "+6", "≈ +65 points"],
        ["Signaux de défausse", "0", "aucun effet mesurable"],
    ], largeurs=[9, 3, 4.5])
    paragraphe(doc, "Ces gains ont été mesurés séparément et ne s'additionnent pas forcément.")
    encadre(doc, "**Prudence** : ces chiffres sont mesurés **dans notre simulation**, face à des "
            "joueurs virtuels. Ils indiquent des tendances solides, mais ne remplacent pas "
            "l'expérience de la table (voir la section « Limites »).")

